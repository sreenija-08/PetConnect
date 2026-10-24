# update_toc_and_clean.py
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import re

DOCX_PATH = r"E:\petconnect\PetConnect_Project_Documentation.docx"

def set_run_font(run, font_name="Times New Roman", size_pt=10, bold=False, italic=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)

def set_cell(cell, text, font_size=10, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    set_run_font(run, "Times New Roman", size_pt=font_size, bold=bold, italic=False)

def set_p(p, text, font_size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, color=None):
    p.text = ""
    p.alignment = align
    run = p.add_run(text)
    set_run_font(run, font_name="Times New Roman", size_pt=font_size, bold=bold, italic=italic, color=color)
    return run

def update_document():
    print("Loading document:", DOCX_PATH)
    doc = docx.Document(DOCX_PATH)

    # 1. Update Table 1: TABLE OF CONTENTS
    t1 = doc.tables[1]
    toc_data = [
        ("Section / Chapter", "Title", "Page"),
        ("CHAPTER 1", "INTRODUCTION", "1"),
        ("1.1", "Introduction", "1"),
        ("1.2", "Problem Statement", "1"),
        ("1.3", "Objectives of the Project", "3"),
        ("1.4", "Scope of the Project", "3"),
        ("CHAPTER 2", "SYSTEM OVERVIEW", "4"),
        ("2.1", "System Architecture", "4"),
        ("2.2", "System Flow", "5"),
        ("2.3", "User Roles", "5"),
        ("2.4", "System Modules", "6"),
        ("CHAPTER 3", "TECHNOLOGIES USED", "7"),
        ("3.1", "Java", "7"),
        ("3.2", "Spring Boot", "7"),
        ("3.3", "Frontend Technologies", "7"),
        ("3.4", "RESTful Web Services", "8"),
        ("3.5", "Spring Data JPA", "8"),
        ("3.6", "MySQL Database", "8"),
        ("3.7", "Validation and Exception Handling", "8"),
        ("3.8", "Spring AOP", "8"),
        ("3.9", "Python FastAPI AI Microservice", "9"),
        ("3.10", "Security Headers and Spring Security", "9"),
        ("3.11", "Maven and Testing Tools", "9"),
        ("CHAPTER 4", "DATABASE DESIGN", "10"),
        ("4.1", "Database Tables Overview", "10"),
        ("4.2", "Shelters Table", "10"),
        ("4.3", "Users Table", "11"),
        ("4.4", "Pets Table", "11"),
        ("4.5", "Adoption Requests Table", "11"),
        ("4.6", "Care Bookings Table", "11"),
        ("4.7", "Shelter Documents Table", "12"),
        ("4.8", "Shelter Visit Slots Table", "12"),
        ("4.9", "Shelter Reviews Table", "12"),
        ("4.10", "Database Relationships and Concurrency Control", "13"),
        ("CHAPTER 5", "SYSTEM MODULES", "14"),
        ("5.1", "Shelter Registration and Onboarding", "14"),
        ("5.2", "User and Adopter Profile Management", "14"),
        ("5.3", "Pet Catalog and Search Filtering", "14"),
        ("5.4", "Adoption Application Workflow", "14"),
        ("5.5", "Temporary Pet Care and Boarding", "14"),
        ("5.6", "AI-Assisted Document Verification", "14"),
        ("5.7", "PetMatch AI Recommendation Engine", "15"),
        ("5.8", "Shelter Visit Scheduling", "15"),
        ("5.9", "Shelter Ratings and Reviews", "15"),
        ("5.10", "Audit Logging, Error Handling and Email Alerts", "15"),
        ("CHAPTER 6", "API AND IMPLEMENTATION", "16"),
        ("6.1", "Frontend Implementation", "16"),
        ("6.2", "Backend Implementation", "16"),
        ("6.3", "REST API Operations", "16"),
        ("6.4", "Transaction Management", "18"),
        ("6.5", "Spring AOP Logging", "18"),
        ("6.6", "Validation and Exception Handling", "18"),
        ("6.7", "AI Microservice Integration", "19"),
        ("6.8", "Security Headers, CORS and Access Control", "19"),
        ("6.9", "Testing", "19"),
        ("6.10", "Running the Project", "20"),
        ("CHAPTER 7", "OUTPUT SCREENS", "21"),
        ("7.1", "Landing Page", "21"),
        ("7.2", "Pet Catalog & Search Filters", "22"),
        ("7.3", "Pet Details Profile", "22"),
        ("7.4", "Adoption Application Form", "23"),
        ("7.5", "PetMatch AI Recommendations", "24"),
        ("7.6", "Temporary Pet Care Booking", "24"),
        ("7.7", "User & Shelter Authentication Portal", "25"),
        ("7.8", "Shelter Owner Dashboard", "26"),
        ("7.9", "Rescued Pet Inventory Management", "26"),
        ("7.10", "Adoption Requests Review Queue", "27"),
        ("7.11", "Admin Shelter Verification Board", "27"),
        ("7.12", "AI Document OCR Verification Details", "28"),
        ("7.13", "PetConnect Platform Statistics & AI Verification API Response", "28"),
        ("7.14", "Automated Test and Build Result", "29"),
        ("CHAPTER 8", "CONCLUSION", "30"),
        ("8.1", "Summary", "30"),
        ("8.2", "Future Enhancements", "30"),
        ("", "REFERENCES", "31")
    ]

    print(f"Trimming Table 1 from {len(t1.rows)} to {len(toc_data)} rows...")
    while len(t1.rows) > len(toc_data):
        t1._tbl.remove(t1.rows[-1]._tr)

    for r_idx, (col0, col1, col2) in enumerate(toc_data):
        is_hdr = (r_idx == 0)
        is_ch = col0.startswith("CHAPTER") or col1 in ["CONCLUSION", "REFERENCES"]
        bold_row = is_hdr or is_ch
        set_cell(t1.rows[r_idx].cells[0], col0, font_size=10, bold=bold_row, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell(t1.rows[r_idx].cells[1], col1, font_size=10, bold=bold_row, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell(t1.rows[r_idx].cells[2], col2, font_size=10, bold=bold_row, align=WD_ALIGN_PARAGRAPH.RIGHT)

    # 2. Update Table 2: LIST OF FIGURES
    t2 = doc.tables[2]
    lof_data = [
        ("Figure", "Description", "Page"),
        ("Figure 2.1", "System Architecture", "4"),
        ("Figure 4.1", "Database Design", "10"),
        ("Figure 7.1", "Landing Page", "21"),
        ("Figure 7.2", "Pet Catalog & Search Filters", "22"),
        ("Figure 7.3", "Pet Details Profile", "22"),
        ("Figure 7.4", "Adoption Application Form", "23"),
        ("Figure 7.5", "PetMatch AI Recommendations", "24"),
        ("Figure 7.6", "Temporary Pet Care Booking", "24"),
        ("Figure 7.7", "User & Shelter Authentication Portal", "25"),
        ("Figure 7.8", "Shelter Owner Dashboard", "26"),
        ("Figure 7.9", "Rescued Pet Inventory Management", "26"),
        ("Figure 7.10", "Adoption Requests Review Queue", "27"),
        ("Figure 7.11", "Admin Shelter Verification Board", "27"),
        ("Figure 7.12", "AI Document OCR Verification Details", "28")
    ]

    print(f"Trimming Table 2 from {len(t2.rows)} to {len(lof_data)} rows...")
    while len(t2.rows) > len(lof_data):
        t2._tbl.remove(t2.rows[-1]._tr)

    for r_idx, (col0, col1, col2) in enumerate(lof_data):
        bold_row = (r_idx == 0)
        set_cell(t2.rows[r_idx].cells[0], col0, font_size=10, bold=bold_row, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell(t2.rows[r_idx].cells[1], col1, font_size=10, bold=bold_row, align=WD_ALIGN_PARAGRAPH.LEFT)
        set_cell(t2.rows[r_idx].cells[2], col2, font_size=10, bold=bold_row, align=WD_ALIGN_PARAGRAPH.RIGHT)

    # 3. Clean up captions and text in paragraphs
    for i, p in enumerate(doc.paragraphs):
        t = p.text.strip()
        # Clean up 7.3 description
        if "The detail page of a hospital shows badges and tabs" in t:
            print(f"Fixing paragraph {i}: 7.3 Pet Details description")
            set_p(p, "The pet details page showcases animal photographs, breed, temperament, vaccination history, medical notes, and the direct adoption application action.", font_size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        
        # Clean up Figure 7.6 caption
        elif "Figure 7.6:" in t or t == "Figure 7.6: Staff Sign-In":
            print(f"Fixing paragraph {i}: Figure 7.6 caption")
            set_p(p, "Figure 7.6: Temporary Pet Care Booking", font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

        # Clean up Figure 7.4 caption
        elif "Figure 7.4:Adoption form" in t or "Figure 7.4: Adoption form" in t:
            print(f"Fixing paragraph {i}: Figure 7.4 caption")
            set_p(p, "Figure 7.4: Adoption Application Form", font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

        # Clean up Figure 7.2 caption
        elif t == "Figure 7.2 Pet Catalog":
            print(f"Fixing paragraph {i}: Figure 7.2 caption")
            set_p(p, "Figure 7.2: Pet Catalog & Search Filters", font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

        # Clean up Figure 7.3 caption
        elif t == "Figure 7.3: Pet Detail":
            print(f"Fixing paragraph {i}: Figure 7.3 caption")
            set_p(p, "Figure 7.3: Pet Details Profile", font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

        # Clean up Figure 7.5 caption
        elif t == "Figure 7.5: Pet Match AI":
            print(f"Fixing paragraph {i}: Figure 7.5 caption")
            set_p(p, "Figure 7.5: PetMatch AI Recommendations", font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

        # Clean up Figure 7.7 caption
        elif t == "Figure 7.7: User and Shelter Sign-Up":
            print(f"Fixing paragraph {i}: Figure 7.7 caption")
            set_p(p, "Figure 7.7: User & Shelter Authentication Portal", font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

        # Clean up Figure 7.9 caption
        elif t == "Figure 7.9: Rescued Pet":
            print(f"Fixing paragraph {i}: Figure 7.9 caption")
            set_p(p, "Figure 7.9: Rescued Pet Inventory Management", font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

        # Clean up Figure 7.10 caption
        elif t == "Figure 7.10: Adoption Request":
            print(f"Fixing paragraph {i}: Figure 7.10 caption")
            set_p(p, "Figure 7.10: Adoption Requests Review Queue", font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

        # Clean up Figure 7.12 caption
        elif t == "Figure 7.12: AI Document OCR Verification":
            print(f"Fixing paragraph {i}: Figure 7.12 caption")
            set_p(p, "Figure 7.12: AI Document OCR Verification Details", font_size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)

        # Clean up 7.11 description style if it was set to Heading 2
        elif t.startswith("Displays the administrator audit dashboard reviewing registered shelters"):
            p.style = doc.styles['Normal']
            set_p(p, t, font_size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

        # Clean up 6.10 MySQL/H2 comment
        elif "To use MySQL instead of H2" in t:
            print(f"Fixing paragraph {i}: 6.10 description")
            set_p(p, "The application connects directly to MySQL 8.x configured in application.properties, and the frontend is rendered dynamically via server-side Thymeleaf templates.", font_size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

        # Clean up 6.10 npm run dev
        elif t == "npm run dev":
            set_p(p, "# Service listening on http://127.0.0.1:8000/", font_size=11, align=WD_ALIGN_PARAGRAPH.LEFT)

        elif "# REST: http://localhost:8086/api/central/census" in t:
            set_p(p, "# Health check: http://localhost:8080/actuator/health", font_size=11, align=WD_ALIGN_PARAGRAPH.LEFT)

        elif "Expected: Tests run: 26" in t:
            set_p(p, "# Expected: Tests run: 24, Failures: 0, Errors: 0, Skipped: 0", font_size=11, align=WD_ALIGN_PARAGRAPH.LEFT)

    # Clean up duplicate [1] [1] in references
    for i in range(len(doc.paragraphs)-20, len(doc.paragraphs)):
        p = doc.paragraphs[i]
        t = p.text.strip()
        m = re.match(r'^\[(\d+)\]\s*\[\1\]\s*(.*)$', t)
        if m:
            print(f"Fixing duplicate numbering in ref {m.group(1)}")
            clean_text = m.group(2)
            set_p(p, clean_text, font_size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
        else:
            # If it starts with [1]\t
            m2 = re.match(r'^\[\d+\]\s*(.*)$', t)
            if m2 and p.style.name == 'List Paragraph':
                set_p(p, m2.group(1), font_size=12, align=WD_ALIGN_PARAGRAPH.JUSTIFY)

    print("Saving updated document to:", DOCX_PATH)
    doc.save(DOCX_PATH)
    print("Saved successfully!")

if __name__ == "__main__":
    update_document()
