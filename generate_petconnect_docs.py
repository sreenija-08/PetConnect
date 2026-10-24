# generate_petconnect_docs.py
# Generates PetConnect_Project_Documentation.docx and PetConnect_Project_Documentation.pdf

import os
import sys
import shutil
import zipfile
import docx
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SOURCE_DOCX = r"C:\Users\sreen\Downloads\CureHub_Project_Documentation.docx"
TARGET_DOCX = r"E:\petconnect\PetConnect_Project_Documentation.docx"
TARGET_PDF = r"E:\petconnect\PetConnect_Project_Documentation.pdf"
DB_IMAGE = r"E:\petconnect\Figure_4_1_Database_Design.png"

DOWNLOADS_DOCX = r"C:\Users\sreen\Downloads\PetConnect_Project_Documentation.docx"
DOWNLOADS_PDF = r"C:\Users\sreen\Downloads\PetConnect_Project_Documentation.pdf"

def set_run_font(run, font_name="Times New Roman", size_pt=12, bold=False, italic=False, color=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)

def set_p_text(p, text, font_size=12, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, color=None):
    p.text = ""
    p.alignment = align
    run = p.add_run(text)
    set_run_font(run, "Times New Roman", size_pt=font_size, bold=bold, italic=italic, color=color)
    return run

def set_p_runs(p, runs_spec, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.text = ""
    p.alignment = align
    for text, sz, b, it, col in runs_spec:
        run = p.add_run(text)
        set_run_font(run, "Times New Roman", size_pt=sz, bold=b, italic=it, color=col)

def set_cell_text(cell, text, font_size=11, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    set_run_font(run, "Times New Roman", size_pt=font_size, bold=bold, italic=False)

print("Helper functions defined.")
